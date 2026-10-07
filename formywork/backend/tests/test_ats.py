import io

from docx import Document
from pypdf import PdfReader

from app.ai.adapt import guard, simple_adaptation
from app.ats.analyzer import analyze
from app.ats.export import to_docx, to_pdf
from app.ats.parser import extract_text, parse_resume_text
from app.ats.schema import ResumeData
from app.demo.seed import DEMO_RESUME
from app.models import Offer

from .conftest import load


def demo() -> ResumeData:
    return ResumeData.model_validate(DEMO_RESUME)


def check(report, check_id):
    return next(c for c in report.checks if c.id == check_id)


def test_good_resume_scores_well_and_explains():
    report = analyze(demo())
    assert 70 <= report.score <= 100
    assert sum(c["max"] for c in report.categories) == 100
    assert check(report, "email").status == "ok"
    for c in report.checks:
        if c.status != "ok":
            assert c.fix, f"{c.id} doit proposer une correction"


def test_empty_resume_scores_low_with_fixes():
    report = analyze(ResumeData())
    assert report.score < 25
    assert check(report, "email").status == "error"
    assert check(report, "section_experience").fix


def test_keywords_against_offer():
    offer_text = "Comptable : révision, fiscalité, Cegid, Excel, facturation et rigueur."
    report = analyze(demo(), offer_text)
    assert "Excel" in report.keywords_matched
    assert "Cegid" in report.keywords_missing
    assert "Cegid" in check(report, "mots_cles_offre").fix


def test_personal_info_and_image_pdf_flags():
    data = demo()
    data.summary = "Née le 12/03/1970, mariée, deux enfants."
    data.import_flags = {"image_only": True, "tables": True}
    report = analyze(data)
    assert check(report, "infos_perso").status == "warning"
    assert check(report, "texte_lisible").status == "error"
    assert check(report, "colonnes").status == "warning"


def test_profile_tips_senior():
    tips = analyze(demo(), profile_type="senior").profile_tips
    assert any("15 dernières années" in t for t in tips)


def test_parse_text_resume():
    data = parse_resume_text(load("cv_exemple.txt"))
    assert data.contact.full_name == "Camille Bernard"
    assert data.contact.email == "camille.bernard@example.org"
    assert data.contact.phone.startswith("06")
    assert len(data.experiences) == 2
    assert data.experiences[0].start == "06/2024"
    assert len(data.experiences[0].bullets) == 2
    assert data.experiences[1].company.startswith("Mairie")
    assert "Recrutement" in data.skills
    assert data.languages[0].name == "Anglais"
    assert data.education


def test_extract_text_rejects_fake_pdf():
    import pytest

    from app.ats.parser import ImportErrorFr

    with pytest.raises(ImportErrorFr):
        extract_text("cv.pdf", b"not a pdf")


def test_pdf_export_is_real_text_and_roundtrips():
    pdf = to_pdf(demo())
    reader = PdfReader(io.BytesIO(pdf))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert "Dominique Laurent" in text
    assert "EXPÉRIENCE PROFESSIONNELLE" in text
    # Le PDF exporté se réimporte correctement
    raw, flags = extract_text("cv.pdf", pdf)
    assert not flags.get("image_only")
    parsed = parse_resume_text(raw)
    assert parsed.contact.email == "dominique.laurent@example.org"


def test_docx_export_single_column():
    doc = Document(io.BytesIO(to_docx(demo())))
    assert not doc.tables
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "FORMATION" in text and "Imprimerie Dufour" in text


def test_guard_removes_invented_numbers_and_skills():
    base = demo()
    adapted = base.model_copy(deep=True)
    adapted.experiences[0].bullets = ["Gérer un portefeuille de 950 clients", "Suivre les règlements"]
    adapted.experiences[0].company = "Entreprise inventée"
    adapted.skills = ["Excel", "SAP", "Facturation"]
    adapted.summary = "Experte avec 40 ans d'expérience."
    alerts = guard(base, adapted)
    assert adapted.experiences[0].bullets == ["Suivre les règlements"]
    assert adapted.experiences[0].company == "Imprimerie Dufour"
    assert "SAP" not in adapted.skills and "Sage" in adapted.skills
    assert adapted.summary == base.summary
    assert len(alerts) == 3


def test_simple_adaptation_reorders_without_inventing():
    offer = Offer(
        title="Assistant(e) de facturation",
        company="X",
        description="Facturation, devis, relances. Maîtrise de SAP souhaitée.",
    )
    adapted, changes, points = simple_adaptation(demo(), offer)
    assert adapted.skills[0] in ("Facturation", "Devis")
    assert sorted(adapted.skills) == sorted(demo().skills)
    assert any("SAP" in p for p in points)
    assert guard(demo(), adapted) == []
