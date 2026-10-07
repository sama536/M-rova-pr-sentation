"""Import d'un CV (PDF, Word .docx, texte) et découpage automatique en sections.

Le découpage est une aide : la personne vérifie et corrige ensuite dans l'éditeur.
"""

from __future__ import annotations

import io
import re

from pypdf import PdfReader
from pypdf.errors import PdfReadError

from ..sources.base import simplify
from .schema import Contact, Education, Experience, Language, ResumeData

ALLOWED_EXT = {".pdf", ".docx", ".txt"}


class ImportErrorFr(Exception):
    pass


def extract_text(filename: str, content: bytes) -> tuple[str, dict]:
    """Renvoie (texte, indicateurs de mise en forme). Vérifie le vrai type du fichier, pas seulement l'extension."""
    name = filename.lower()
    flags: dict = {}
    if name.endswith(".pdf"):
        if not content.startswith(b"%PDF"):
            raise ImportErrorFr("Ce fichier n'est pas un vrai PDF.")
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ImportErrorFr("Ce PDF est protégé par un mot de passe.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:6])
        except PdfReadError as exc:
            raise ImportErrorFr("Impossible de lire ce PDF (fichier abîmé ?).") from exc
        if len(text.strip()) < 80:
            flags["image_only"] = True
        if _looks_multicolumn(text):
            flags["columns"] = True
        return text, flags
    if name.endswith(".docx"):
        if not content.startswith(b"PK"):
            raise ImportErrorFr("Ce fichier n'est pas un vrai document Word (.docx).")
        from docx import Document

        try:
            doc = Document(io.BytesIO(content))
        except Exception as exc:  # python-docx lève des erreurs variées sur un fichier invalide
            raise ImportErrorFr("Impossible de lire ce document Word.") from exc
        lines = [p.text for p in doc.paragraphs]
        if doc.tables:
            flags["tables"] = True
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        lines.extend(cell.text.splitlines())
        return "\n".join(lines), flags
    if name.endswith(".txt"):
        for enc in ("utf-8", "cp1252", "latin-1"):
            try:
                return content.decode(enc), flags
            except UnicodeDecodeError:
                continue
    raise ImportErrorFr("Format non pris en charge. Utilisez un fichier PDF, Word (.docx) ou texte (.txt).")


def _looks_multicolumn(text: str) -> bool:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 15:
        return False
    wide_gaps = sum(1 for ln in lines if re.search(r"\S\s{4,}\S", ln))
    return wide_gaps / len(lines) > 0.25


SECTION_KEYS = {
    "summary": ["profil", "a propos", "resume", "accroche", "presentation", "objectif"],
    "experiences": ["experience", "experiences professionnelles", "parcours professionnel", "emplois", "carriere"],
    "education": ["formation", "formations", "diplome", "diplomes", "etudes", "scolarite", "cursus"],
    "skills": ["competences", "savoir faire", "outils", "informatique", "aptitudes", "qualites"],
    "languages": ["langues", "langue"],
    "certifications": ["certifications", "certificats", "habilitations", "permis"],
    "interests": ["centres d interet", "loisirs", "interets", "activites", "benevolat"],
}
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE = re.compile(r"(?:\+33\s?|0)[1-9](?:[\s.-]?\d{2}){4}")
_DATE_RANGE = re.compile(
    r"((?:0?[1-9]|1[0-2])[/.-](?:19|20)\d{2}|(?:janv|fevr|févr|mars|avr|mai|juin|juil|aout|août|sept|oct|nov|dec|déc)\w*\.?\s+(?:19|20)\d{2}|(?:19|20)\d{2})"
    r"\s*(?:-|–|—|à|au|a)\s*"
    r"((?:0?[1-9]|1[0-2])[/.-](?:19|20)\d{2}|(?:janv|fevr|févr|mars|avr|mai|juin|juil|aout|août|sept|oct|nov|dec|déc)\w*\.?\s+(?:19|20)\d{2}|(?:19|20)\d{2}|aujourd'hui|aujourd’hui|present|présent|en cours|maintenant)",
    re.I,
)
_BULLET = re.compile(r"^\s*[•●▪◦\-–*·]\s*")
_MONTHS = {
    "jan": "01",
    "fev": "02",
    "mar": "03",
    "avr": "04",
    "mai": "05",
    "juin": "06",
    "juil": "07",
    "aou": "08",
    "sep": "09",
    "oct": "10",
    "nov": "11",
    "dec": "12",
}


def _section_of(line: str) -> str | None:
    s = simplify(line)
    if not s or len(s) > 40:
        return None
    for key, words in SECTION_KEYS.items():
        if any(s == w or s.startswith(w) for w in words):
            return key
    return None


def _norm_date(raw: str) -> str:
    raw = raw.strip()
    s = simplify(raw)
    if any(w in s for w in ("aujourd", "present", "cours", "maintenant")):
        return ""
    m = re.match(r"(\d{1,2})\D(\d{4})", raw)
    if m:
        return f"{int(m.group(1)):02d}/{m.group(2)}"
    year = re.search(r"(19|20)\d{2}", raw)
    for prefix, num in _MONTHS.items():
        if s.startswith(prefix) and year:
            return f"{num}/{year.group(0)}"
    return year.group(0) if year else raw


def parse_resume_text(text: str) -> ResumeData:
    lines = [ln.rstrip() for ln in text.replace("\r", "").splitlines()]
    data = ResumeData(contact=Contact())
    email = _EMAIL.search(text)
    phone = _PHONE.search(text)
    data.contact.email = email.group(0) if email else ""
    data.contact.phone = phone.group(0) if phone else ""
    data.contact.links = sorted(set(re.findall(r"(?:https?://|www\.)\S+|linkedin\.com/\S+", text)))[:3]

    header: list[str] = []
    sections: dict[str, list[str]] = {k: [] for k in SECTION_KEYS}
    current: str | None = None
    for line in lines:
        if not line.strip():
            if current:
                sections[current].append("")
            continue
        sec = _section_of(line)
        if sec:
            current = sec
            continue
        if current is None:
            header.append(line.strip())
        else:
            sections[current].append(line.strip())

    clean_header = [h for h in header if not _EMAIL.search(h) and not _PHONE.search(h) and "http" not in h]
    if clean_header:
        data.contact.full_name = clean_header[0][:80]
    if len(clean_header) > 1:
        data.contact.headline = clean_header[1][:120]
    for h in clean_header[1:]:
        if re.search(r"\b\d{5}\b", h) or (len(h.split()) <= 3 and h[:1].isupper() and h != data.contact.headline):
            data.contact.city = re.sub(r"\b\d{5}\b", "", h).strip(" ,-")
            break

    data.summary = " ".join(x for x in sections["summary"] if x).strip()
    data.experiences = _parse_items(sections["experiences"], Experience)
    data.education = [
        Education(degree=e.title, school=e.company, start=e.start, end=e.end, details=" ".join(e.bullets)[:300])
        for e in _parse_items(sections["education"], Experience)
    ]
    data.skills = _split_list(sections["skills"])
    data.languages = [
        Language(
            name=part.split(":")[0].split("(")[0].strip(), level=(part.split(":", 1)[1] if ":" in part else "").strip()
        )
        for part in _split_list(sections["languages"])
    ]
    data.certifications = _split_list(sections["certifications"])
    data.interests = _split_list(sections["interests"])
    return data


def _split_list(lines: list[str]) -> list[str]:
    items: list[str] = []
    for line in lines:
        line = _BULLET.sub("", line)
        for part in re.split(r"\s*[,;|•]\s*", line):
            part = part.strip(" .")
            if 1 < len(part) <= 80 and part not in items:
                items.append(part)
    return items[:30]


def _parse_items(lines: list[str], cls):
    items = []
    cur = None
    pending_title: list[str] = []
    for line in lines:
        if not line:
            continue
        dates = _DATE_RANGE.search(line)
        if dates:
            rest = (line[: dates.start()] + " " + line[dates.end() :]).strip(" |,-–—:")
            title_parts = [p for p in [*pending_title, rest] if p]
            cur = cls(start=_norm_date(dates.group(1)), end=_norm_date(dates.group(2)))
            cur.current = not cur.end
            if title_parts:
                head = title_parts[0]
                if " - " in head or " – " in head or "," in head:
                    t, _, comp = re.split(r"(\s[-–]\s|,)", head, maxsplit=1)
                    cur.title, cur.company = t.strip(), comp.strip()
                else:
                    cur.title = head
                    if len(title_parts) > 1:
                        cur.company = title_parts[1]
            pending_title = []
            items.append(cur)
        elif _BULLET.match(line) and cur is not None:
            cur.bullets.append(_BULLET.sub("", line))
        elif cur is not None and not cur.company and not cur.bullets:
            cur.company = line
        elif cur is not None and len(line) > 60:
            cur.bullets.append(line)
        else:
            pending_title.append(line)
            if len(pending_title) > 2:
                pending_title = pending_title[-2:]
    return items
