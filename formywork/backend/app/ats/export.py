"""Exports PDF (vrai texte, sélectionnable) et DOCX, modèle sobre en une colonne, titres standard."""

from __future__ import annotations

import io
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from .schema import ResumeData

NAVY = "#14264A"

HEADINGS = {
    "summary": "PROFIL",
    "experiences": "EXPÉRIENCE PROFESSIONNELLE",
    "education": "FORMATION",
    "skills": "COMPÉTENCES",
    "languages": "LANGUES",
    "certifications": "CERTIFICATIONS",
    "interests": "CENTRES D'INTÉRÊT",
}


def _period(start: str, end: str, current: bool) -> str:
    if not start and not end:
        return ""
    return f"{start or '?'} – {'aujourd’hui' if current or not end else end}"


def contact_line(data: ResumeData) -> str:
    c = data.contact
    return " | ".join(p for p in [c.city, c.phone, c.email, *c.links] if p)


def to_pdf(data: ResumeData) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"CV {data.contact.full_name}".strip(),
        author=data.contact.full_name or "FormyWork",
    )
    navy = HexColor(NAVY)
    st = {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=navy),
        "headline": ParagraphStyle("headline", fontName="Helvetica", fontSize=11.5, leading=15, textColor=navy),
        "contact": ParagraphStyle("contact", fontName="Helvetica", fontSize=9.5, leading=13),
        "h": ParagraphStyle(
            "h", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=navy, spaceBefore=9, spaceAfter=3
        ),
        "item": ParagraphStyle("item", fontName="Helvetica-Bold", fontSize=10.5, leading=13.5),
        "meta": ParagraphStyle("meta", fontName="Helvetica-Oblique", fontSize=9.5, leading=12.5),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=13.5, alignment=TA_LEFT),
        "bullet": ParagraphStyle(
            "bullet", fontName="Helvetica", fontSize=10, leading=13.5, leftIndent=10, bulletIndent=0
        ),
    }
    e = lambda s: escape(s or "")  # noqa: E731
    flow: list = [Paragraph(e(data.contact.full_name) or "Votre nom", st["name"])]
    if data.contact.headline:
        flow.append(Paragraph(e(data.contact.headline), st["headline"]))
    if contact_line(data):
        flow.append(Paragraph(e(contact_line(data)), st["contact"]))
    flow.append(Spacer(1, 4))

    def heading(key: str) -> None:
        flow.append(Paragraph(HEADINGS[key], st["h"]))
        flow.append(HRFlowable(width="100%", thickness=0.6, color=navy, spaceAfter=4))

    if data.summary:
        heading("summary")
        flow.append(Paragraph(e(data.summary), st["body"]))
    if data.experiences:
        heading("experiences")
        for x in data.experiences:
            flow.append(Paragraph(e(" – ".join(p for p in [x.title, x.company] if p)), st["item"]))
            meta = " | ".join(p for p in [_period(x.start, x.end, x.current), x.location] if p)
            if meta:
                flow.append(Paragraph(e(meta), st["meta"]))
            for bullet in x.bullets:
                if bullet.strip():
                    flow.append(Paragraph(e(bullet), st["bullet"], bulletText="•"))
            flow.append(Spacer(1, 5))
    if data.education:
        heading("education")
        for ed in data.education:
            flow.append(Paragraph(e(" – ".join(p for p in [ed.degree, ed.school] if p)), st["item"]))
            meta = " | ".join(p for p in [_period(ed.start, ed.end, False) if ed.start else ed.end, ed.location] if p)
            if meta:
                flow.append(Paragraph(e(meta), st["meta"]))
            if ed.details:
                flow.append(Paragraph(e(ed.details), st["body"]))
            flow.append(Spacer(1, 4))
    if data.skills:
        heading("skills")
        flow.append(Paragraph(e(", ".join(data.skills)), st["body"]))
    if data.languages:
        heading("languages")
        flow.append(
            Paragraph(
                e(", ".join(f"{lg.name} : {lg.level}" if lg.level else lg.name for lg in data.languages)), st["body"]
            )
        )
    if data.certifications:
        heading("certifications")
        for cert in data.certifications:
            flow.append(Paragraph(e(cert), st["bullet"], bulletText="•"))
    if data.interests:
        heading("interests")
        flow.append(Paragraph(e(", ".join(data.interests)), st["body"]))
    doc.build(flow)
    return buf.getvalue()


def to_docx(data: ResumeData) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    navy = RGBColor(0x14, 0x26, 0x4A)

    def para(text: str, *, bold=False, italic=False, size=None, color=None, space_after=2):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def heading(key: str):
        h = doc.add_heading(HEADINGS[key], level=2)
        for run in h.runs:
            run.font.color.rgb = navy
            run.font.size = Pt(12)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title = doc.add_heading(data.contact.full_name or "Votre nom", level=1)
    for run in title.runs:
        run.font.color.rgb = navy
    if data.contact.headline:
        para(data.contact.headline, size=12, color=navy)
    if contact_line(data):
        para(contact_line(data), space_after=8)
    if data.summary:
        heading("summary")
        para(data.summary)
    if data.experiences:
        heading("experiences")
        for x in data.experiences:
            para(" – ".join(p for p in [x.title, x.company] if p), bold=True, space_after=0)
            meta = " | ".join(p for p in [_period(x.start, x.end, x.current), x.location] if p)
            if meta:
                para(meta, italic=True, space_after=1)
            for bullet in x.bullets:
                if bullet.strip():
                    doc.add_paragraph(bullet, style="List Bullet")
    if data.education:
        heading("education")
        for ed in data.education:
            para(" – ".join(p for p in [ed.degree, ed.school] if p), bold=True, space_after=0)
            meta = " | ".join(p for p in [_period(ed.start, ed.end, False) if ed.start else ed.end, ed.location] if p)
            if meta:
                para(meta, italic=True, space_after=1)
            if ed.details:
                para(ed.details)
    if data.skills:
        heading("skills")
        para(", ".join(data.skills))
    if data.languages:
        heading("languages")
        para(", ".join(f"{lg.name} : {lg.level}" if lg.level else lg.name for lg in data.languages))
    if data.certifications:
        heading("certifications")
        for cert in data.certifications:
            doc.add_paragraph(cert, style="List Bullet")
    if data.interests:
        heading("interests")
        para(", ".join(data.interests))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def letter_to_pdf(text: str, author: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=22 * mm,
        bottomMargin=20 * mm,
        title="Lettre de motivation",
        author=author,
    )
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=11, leading=15.5, spaceAfter=9)
    flow = [Paragraph(escape(p).replace("\n", "<br/>"), body) for p in text.split("\n\n") if p.strip()]
    doc.build(flow)
    return buf.getvalue()
